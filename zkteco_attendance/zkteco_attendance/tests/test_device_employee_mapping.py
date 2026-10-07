"""
Unit tests for the Biometric Device "Browse Employees On Device" feature:
- get_device_users endpoint (fetch + annotate device users)
- map_device_employees endpoint (persist attendance_device_id mapping)
- shift_type handling in both endpoints (ZK Shift Assignment updates)
- download_device_users_excel endpoint (dialog's Download Excel button)
Run with: bench run-tests --app zkteco_attendance
"""

import base64
import json
import unittest
from io import BytesIO
from unittest.mock import patch, MagicMock, call

import frappe
from openpyxl import load_workbook


class _FakeTemplate:
    """Minimal stand-in for a pyzk Finger template object."""

    def __init__(self, uid, valid=1):
        self.uid = uid
        self.valid = valid


class TestGetDeviceUsersEndpoint(unittest.TestCase):
    """get_device_users annotates device users with their mapped Employee."""

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_zk_connection")
    @patch("frappe.get_all")
    def test_annotates_users_with_mapped_employee(self, mock_get_all, mock_conn_fn, mock_reachable):
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        mock_conn = MagicMock()
        user = MagicMock(uid=1, user_id="100", name="Abebe Kebede", privilege=0)
        mock_conn.get_users.return_value = [user]
        mock_conn_fn.return_value = (mock_conn, MagicMock())

        # frappe.get_all returns frappe._dict rows (attribute access).
        mock_get_all.return_value = [
            frappe._dict({"name": "HR-EMP-00001", "employee_name": "Abebe Kebede",
                          "attendance_device_id": "100"})
        ]

        with patch("frappe.get_doc", return_value=MagicMock()):
            result = get_device_users("Test-ZK-Device")

        self.assertTrue(result["success"])
        self.assertEqual(result["count"], 1)
        user_row = result["users"][0]
        self.assertEqual(user_row["user_id"], "100")
        self.assertEqual(user_row["employee"], "HR-EMP-00001")
        self.assertEqual(user_row["employee_name"], "Abebe Kebede")

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_zk_connection")
    @patch("frappe.get_all")
    def test_unmapped_users_have_empty_employee(self, mock_get_all, mock_conn_fn, mock_reachable):
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        mock_conn = MagicMock()
        mock_conn.get_users.return_value = [
            MagicMock(uid=2, user_id="200", name="Unmapped User", privilege=0)
        ]
        mock_conn_fn.return_value = (mock_conn, MagicMock())
        mock_get_all.return_value = []

        with patch("frappe.get_doc", return_value=MagicMock()):
            result = get_device_users("Test-ZK-Device")

        self.assertEqual(result["users"][0]["employee"], "")
        self.assertEqual(result["users"][0]["employee_name"], "")
        self.assertEqual(result["users"][0]["shift_type"], "")

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_zk_connection")
    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._get_active_shift_types")
    @patch("frappe.get_all")
    def test_mapped_user_annotated_with_current_shift_type(
            self, mock_get_all, mock_shifts, mock_conn_fn, mock_reachable):
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        mock_conn = MagicMock()
        mock_conn.get_users.return_value = [
            MagicMock(uid=1, user_id="100", name="Abebe Kebede", privilege=0)
        ]
        mock_conn_fn.return_value = (mock_conn, MagicMock())

        mock_get_all.return_value = [
            frappe._dict({"name": "HR-EMP-00001", "employee_name": "Abebe Kebede",
                          "attendance_device_id": "100"})
        ]
        mock_shifts.return_value = {
            "HR-EMP-00001": {"assignment": "ZK-SA-0001", "shift_type": "Morning"},
        }

        with patch("frappe.get_doc", return_value=MagicMock()):
            result = get_device_users("Test-ZK-Device")

        self.assertEqual(result["users"][0]["shift_type"], "Morning")

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_zk_connection")
    @patch("frappe.get_all", return_value=[])
    def test_users_carry_enrolled_punch_methods(self, mock_get_all, mock_conn_fn, mock_reachable):
        """Fingerprint templates, password and card are summarised per user."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        mock_conn = MagicMock()
        u1 = MagicMock(uid=1, user_id="100", name="A", privilege=0)
        u1.password = "1234"
        u1.card = 0
        u2 = MagicMock(uid=2, user_id="200", name="B", privilege=0)
        u2.password = ""
        u2.card = 98765432
        mock_conn.get_users.return_value = [u1, u2]
        # uid 1 has two valid templates; uid 2 only an empty (invalid) slot.
        mock_conn.get_templates.return_value = [
            _FakeTemplate(1), _FakeTemplate(1), _FakeTemplate(2, valid=0),
        ]
        mock_conn_fn.return_value = (mock_conn, MagicMock())

        with patch("frappe.get_doc", return_value=MagicMock()):
            result = get_device_users("Test-ZK-Device")

        self.assertEqual(result["users"][0]["punch_methods"], ["Fingerprint", "Password"])
        self.assertEqual(result["users"][1]["punch_methods"], ["Card"])

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_zk_connection")
    @patch("frappe.get_all", return_value=[])
    def test_punch_methods_empty_when_nothing_enrolled(self, mock_get_all, mock_conn_fn, mock_reachable):
        """A user with no fingerprint/password/card reports an empty list."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        mock_conn = MagicMock()
        u = MagicMock(uid=5, user_id="500", name="C", privilege=0)
        u.password = ""
        u.card = 0
        mock_conn.get_users.return_value = [u]
        mock_conn.get_templates.return_value = []
        mock_conn_fn.return_value = (mock_conn, MagicMock())

        with patch("frappe.get_doc", return_value=MagicMock()):
            result = get_device_users("Test-ZK-Device")

        self.assertEqual(result["users"][0]["punch_methods"], [])

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_zk_connection")
    @patch("frappe.get_all", return_value=[])
    def test_template_read_failure_still_returns_users(self, mock_get_all, mock_conn_fn, mock_reachable):
        """A failed template read leaves the fingerprint badge off but the
        dialog still gets the users."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        mock_conn = MagicMock()
        u = MagicMock(uid=1, user_id="100", name="A", privilege=0)
        u.password = ""
        u.card = 0
        mock_conn.get_users.return_value = [u]
        mock_conn.get_templates.side_effect = Exception("unsupported firmware")
        mock_conn_fn.return_value = (mock_conn, MagicMock())

        with patch("frappe.get_doc", return_value=MagicMock()), \
                patch("frappe.log_error"):
            result = get_device_users("Test-ZK-Device")

        self.assertEqual(result["count"], 1)
        self.assertEqual(result["users"][0]["punch_methods"], [])

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(False, "Device cannot be reached. Check that it is powered on and connected to the network, then try again."))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_device_users")
    def test_unreachable_device_returns_structured_error(self, mock_fetch, mock_reachable):
        """An offline device yields success=False + error so the dialog can
        show a message instead of raising (which the JS callback cannot see)."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        with patch("frappe.get_doc", return_value=MagicMock()):
            result = get_device_users("Test-ZK-Device")

        self.assertFalse(result["success"])
        self.assertIn("cannot be reached", result["error"])
        self.assertEqual(result["users"], [])
        # The device must not be contacted when the probe fails.
        mock_fetch.assert_not_called()

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.get_device_users",
           side_effect=frappe.ValidationError("Cannot connect to device Test-ZK-Device at 192.0.2.55:4370. Error: timed out"))
    def test_fetch_error_returns_structured_error(self, mock_fetch, mock_reachable):
        """A connection error raised while fetching is returned as data, not thrown."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import get_device_users

        device = MagicMock()
        device.device_name = "Test-ZK-Device"
        with patch("frappe.get_doc", return_value=device):
            result = get_device_users("Test-ZK-Device")

        self.assertFalse(result["success"])
        self.assertIn("Could not fetch users", result["error"])
        self.assertEqual(result["count"], 0)


class TestMapDeviceEmployeesEndpoint(unittest.TestCase):
    """map_device_employees persists / clears device-to-employee mappings."""

    def _device(self, name="Test-ZK-Device", company="Acme", ignore_company=0):
        device = MagicMock()
        device.name = name
        device.company = company
        device.ignore_company_restriction = ignore_company
        return device

    def test_rejects_non_list_mappings(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        with patch("frappe.get_doc", return_value=self._device()), \
                self.assertRaises(frappe.ValidationError):
            map_device_employees("Test-ZK-Device", mappings="not-a-list")

    def test_rejects_invalid_json(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        with patch("frappe.get_doc", return_value=self._device()), \
                self.assertRaises(frappe.ValidationError):
            map_device_employees("Test-ZK-Device", mappings="{bad json")

    def test_rejects_mapping_without_user_id(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        with patch("frappe.get_doc", return_value=self._device()), \
                self.assertRaises(frappe.ValidationError):
            map_device_employees("Test-ZK-Device", mappings=[{"employee": "HR-EMP-00001"}])

    def test_rejects_duplicate_user_ids(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [
            {"user_id": "100", "employee": "HR-EMP-00001"},
            {"user_id": "100", "employee": "HR-EMP-00002"},
        ]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Acme"), \
                self.assertRaises(frappe.ValidationError):
            map_device_employees("Test-ZK-Device", mappings=mappings)

    def test_rejects_nonexistent_employee(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=False), \
                self.assertRaises(frappe.ValidationError):
            map_device_employees(
                "Test-ZK-Device", mappings=[{"user_id": "100", "employee": "HR-EMP-404"}])

    def test_rejects_cross_company_employee(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        with patch("frappe.get_doc", return_value=self._device(company="Acme")), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Other Co"), \
                self.assertRaises(frappe.ValidationError):
            map_device_employees(
                "Test-ZK-Device", mappings=[{"user_id": "100", "employee": "HR-EMP-00001"}])

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints.frappe.db.set_value")
    def test_allows_cross_company_employee_when_ignore_restriction(
            self, mock_set_value):
        """With 'Ignore Company Restriction' on, employees of any company map."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        with patch("frappe.get_doc", return_value=self._device(ignore_company=1)), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Other Co"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.commit"):
            result = map_device_employees(
                "Test-ZK-Device", mappings=[{"user_id": "100", "employee": "HR-EMP-00001"}])

        self.assertTrue(result["success"])
        self.assertEqual(result["mapped"], 1)

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints.frappe.db.set_value")
    def test_allows_any_employee_when_device_has_no_company(self, mock_set_value):
        """A device without a company (ignore flag on) accepts any employee."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        with patch("frappe.get_doc", return_value=self._device(company=None, ignore_company=1)), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Other Co"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.commit"):
            result = map_device_employees(
                "Test-ZK-Device", mappings=[{"user_id": "100", "employee": "HR-EMP-00001"}])

        self.assertTrue(result["success"])
        self.assertEqual(result["mapped"], 1)

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints.frappe.db.set_value")
    def test_sets_device_id_and_device_on_employee(self, mock_set_value):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001"}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Acme"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.commit") as mock_commit:
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        self.assertTrue(result["success"])
        self.assertEqual(result["mapped"], 1)
        self.assertEqual(result["unmapped"], 0)

        args, kwargs = mock_set_value.call_args
        self.assertEqual(args[0], "Employee")
        self.assertEqual(args[1], "HR-EMP-00001")
        self.assertEqual(args[2], {
            "attendance_device_id": "100",
            "zk_biometric_device": "Test-ZK-Device",
        })
        mock_commit.assert_called_once()

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints.frappe.db.set_value")
    def test_empty_employee_clears_mapping(self, mock_set_value):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": ""}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        self.assertTrue(result["success"])
        self.assertEqual(result["mapped"], 0)
        self.assertEqual(result["unmapped"], 1)
        mock_set_value.assert_not_called()

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints.frappe.db.set_value")
    def test_stale_mapping_is_cleared(self, mock_set_value):
        """
        An Employee previously mapped to a device user_id that is no longer
        mapped in the submitted payload gets its attendance_device_id cleared.
        """
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001"}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Acme"), \
                patch("frappe.get_all",
                      return_value=[frappe._dict({"name": "HR-EMP-00009"})]), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        self.assertTrue(result["success"])
        # One stale-clear (None) and one mapping (dict) write
        calls = mock_set_value.call_args_list
        self.assertEqual(len(calls), 2)
        self.assertEqual(calls[0].args[1], "HR-EMP-00009")
        # set_value(doctype, name, key, value) — the stale clear writes None
        self.assertEqual(calls[0].args[2], "attendance_device_id")
        self.assertIsNone(calls[0].args[3])
        self.assertEqual(calls[1].args[1], "HR-EMP-00001")
        self.assertEqual(calls[1].args[2]["attendance_device_id"], "100")

    def test_accepts_json_string_payload(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = json.dumps([{"user_id": "100", "employee": "HR-EMP-00001"}])
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Acme"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        self.assertTrue(result["success"])
        self.assertEqual(result["mapped"], 1)


class TestMapDeviceEmployeesShifts(unittest.TestCase):
    """shift_type in mappings drives ZK Shift Assignment changes."""

    def _device(self, name="Test-ZK-Device", company="Acme", ignore_company=0):
        device = MagicMock()
        device.name = name
        device.company = company
        device.ignore_company_restriction = ignore_company
        return device

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._apply_shift_assignment")
    def test_shift_type_passed_to_apply(self, mock_apply):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001",
                     "shift_type": "Morning"}]

        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Acme"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        mock_apply.assert_called_once_with("HR-EMP-00001", "Morning", "Acme")
        self.assertEqual(result["shifts_assigned"], 1)

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._apply_shift_assignment",
           return_value=False)
    def test_unchanged_shift_not_counted(self, mock_apply):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001",
                     "shift_type": "Morning"}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Acme"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        self.assertEqual(result["shifts_assigned"], 0)

    def test_unknown_shift_type_rejected(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001",
                     "shift_type": "Ghost Shift"}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists") as mock_exists, \
                patch("frappe.db.get_value", return_value="Acme"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"), \
                self.assertRaises(frappe.ValidationError):
            # Employee exists, but the ZK Shift Type does not.
            mock_exists.side_effect = lambda doctype, name=None: doctype == "Employee"
            map_device_employees("Test-ZK-Device", mappings=mappings)

        # The validation failed on the unknown ZK Shift Type.
        self.assertIn(call("ZK Shift Type", "Ghost Shift"), mock_exists.call_args_list)

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._apply_shift_assignment")
    def test_allows_cross_company_shift_when_device_ignores_restriction(
            self, mock_apply):
        """With 'Ignore Company Restriction' on, any company's shift is allowed."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001",
                     "shift_type": "Night"}]
        with patch("frappe.get_doc", return_value=self._device(ignore_company=1)), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Other Co"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        mock_apply.assert_called_once_with("HR-EMP-00001", "Night", "Acme")
        self.assertEqual(result["shifts_assigned"], 1)

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._apply_shift_assignment")
    def test_allows_cross_company_shift_when_shift_ignores_restriction(
            self, mock_apply):
        """A ZK Shift Type that ignores company restrictions is usable by any device."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        def fake_get_value(doctype, name, field):
            if doctype == "Employee":
                return "Acme"       # employee matches the device company
            if field == "company":
                return "Other Co"   # shift type belongs to another company
            return 1                # ignore_company_restriction on the shift type

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001",
                     "shift_type": "Night"}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", side_effect=fake_get_value), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        mock_apply.assert_called_once_with("HR-EMP-00001", "Night", "Acme")

    def test_rejects_cross_company_shift_when_no_ignore(self):
        """Without any ignore flag, a cross-company shift type is still rejected."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        def fake_get_value(doctype, name, field):
            if doctype == "Employee":
                return "Acme"       # employee matches the device company
            if field == "company":
                return "Other Co"   # shift type belongs to another company
            return 0                # ignore_company_restriction off

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001",
                     "shift_type": "Night"}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", side_effect=fake_get_value), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"), \
                self.assertRaises(frappe.ValidationError):
            map_device_employees("Test-ZK-Device", mappings=mappings)

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._apply_shift_assignment")
    def test_empty_shift_still_passed_for_unassign(self, mock_apply):
        """An explicit empty shift_type key unassigns the employee."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import map_device_employees

        mappings = [{"user_id": "100", "employee": "HR-EMP-00001", "shift_type": ""}]
        with patch("frappe.get_doc", return_value=self._device()), \
                patch("frappe.db.exists", return_value=True), \
                patch("frappe.db.get_value", return_value="Acme"), \
                patch("frappe.get_all", return_value=[]), \
                patch("frappe.db.set_value"), \
                patch("frappe.db.commit"):
            result = map_device_employees("Test-ZK-Device", mappings=mappings)

        mock_apply.assert_called_once_with("HR-EMP-00001", "", "Acme")
        self.assertEqual(result["shifts_assigned"], 1)



class TestApplyShiftAssignmentCompany(unittest.TestCase):
    """_apply_shift_assignment handles company-less (ignore-restriction) devices."""

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._get_active_shift_types",
           return_value={})
    def test_new_assignment_without_company_sets_ignore_flag(self, mock_shifts):
        from zkteco_attendance.zkteco_attendance.api.endpoints import _apply_shift_assignment

        sa_doc = MagicMock()
        with patch("frappe.get_all", return_value=[]), \
                patch("frappe.new_doc", return_value=sa_doc), \
                patch("frappe.db.get_value", return_value={}), \
                patch("frappe.get_doc"):
            changed = _apply_shift_assignment("HR-EMP-00001", "Morning", None)

        self.assertTrue(changed)
        self.assertIsNone(sa_doc.company)
        self.assertEqual(sa_doc.ignore_company_restriction, 1)
        sa_doc.save.assert_called_once()

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._get_active_shift_types",
           return_value={})
    def test_new_assignment_with_company_keeps_company(self, mock_shifts):
        from zkteco_attendance.zkteco_attendance.api.endpoints import _apply_shift_assignment

        sa_doc = MagicMock()
        sa_doc.ignore_company_restriction = 0
        with patch("frappe.get_all", return_value=[]), \
                patch("frappe.new_doc", return_value=sa_doc), \
                patch("frappe.db.get_value", return_value={}), \
                patch("frappe.get_doc"):
            changed = _apply_shift_assignment("HR-EMP-00001", "Morning", "Acme")

        self.assertTrue(changed)
        self.assertEqual(sa_doc.company, "Acme")
        self.assertEqual(sa_doc.ignore_company_restriction, 0)

    @patch("zkteco_attendance.zkteco_attendance.api.endpoints._get_active_shift_types",
           return_value={})
    def test_matches_assignment_of_same_empty_company(self, mock_shifts):
        """A company-less device joins the company-less assignment of the shift."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import _apply_shift_assignment

        existing = MagicMock()
        with patch("frappe.get_all", return_value=[
                frappe._dict({"name": "ZK-SA-0001", "company": "Acme"}),
                frappe._dict({"name": "ZK-SA-0002", "company": None}),
            ]), \
                patch("frappe.get_doc", return_value=existing) as mock_get_doc, \
                patch("frappe.new_doc") as mock_new_doc:
            changed = _apply_shift_assignment("HR-EMP-00001", "Morning", None)

        self.assertTrue(changed)
        mock_get_doc.assert_called_once_with("ZK Shift Assignment", "ZK-SA-0002")
        mock_new_doc.assert_not_called()
        existing.save.assert_called_once()


class TestDeleteDeviceUserEndpoint(unittest.TestCase):
    """delete_device_user removes a device user and clears its mapping."""

    def _device(self, name="Test-ZK-Device"):
        device = MagicMock()
        device.name = name
        device.device_name = "Test ZK Device"
        return device

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.delete_device_user",
           return_value=True)
    @patch("frappe.get_all", return_value=[])
    def test_deletes_user_and_reports_success(self, mock_get_all, mock_delete, mock_reachable):
        from zkteco_attendance.zkteco_attendance.api.endpoints import delete_device_user

        with patch("frappe.get_doc", return_value=self._device()):
            result = delete_device_user("Test-ZK-Device", "100")

        self.assertTrue(result["success"])
        self.assertEqual(result["deleted"], "100")
        self.assertEqual(result["unmapped"], 0)
        mock_delete.assert_called_once_with("Test-ZK-Device", "100")

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(False, "Device cannot be reached. Check that it is powered on and connected to the network, then try again."))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.delete_device_user")
    def test_unreachable_device_returns_structured_error(self, mock_delete, mock_reachable):
        """An offline device yields success=False + error without contacting it."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import delete_device_user

        with patch("frappe.get_doc", return_value=self._device()):
            result = delete_device_user("Test-ZK-Device", "100")

        self.assertFalse(result["success"])
        self.assertIn("cannot be reached", result["error"])
        mock_delete.assert_not_called()

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.delete_device_user",
           side_effect=frappe.ValidationError("User 100 was not found on device Test ZK Device."))
    def test_delete_failure_returns_structured_error(self, mock_delete, mock_reachable):
        """A device-side failure is returned as data, not raised."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import delete_device_user

        with patch("frappe.get_doc", return_value=self._device()):
            result = delete_device_user("Test-ZK-Device", "100")

        self.assertFalse(result["success"])
        self.assertIn("Could not delete user 100", result["error"])

    @patch("zkteco_attendance.zkteco_attendance.zk_client.check_device_reachable",
           return_value=(True, ""))
    @patch("zkteco_attendance.zkteco_attendance.zk_client.delete_device_user",
           return_value=True)
    @patch("frappe.db.commit")
    @patch("frappe.db.set_value")
    @patch("frappe.get_all",
           return_value=[frappe._dict({"name": "HR-EMP-00001"})])
    def test_clears_mapped_employee(self, mock_get_all, mock_set_value, mock_commit,
                                    mock_delete, mock_reachable):
        """A mapped Employee's attendance_device_id is cleared on delete."""
        from zkteco_attendance.zkteco_attendance.api.endpoints import delete_device_user

        with patch("frappe.get_doc", return_value=self._device()):
            result = delete_device_user("Test-ZK-Device", "100")

        self.assertTrue(result["success"])
        self.assertEqual(result["unmapped"], 1)
        mock_set_value.assert_called_once_with(
            "Employee", "HR-EMP-00001", "attendance_device_id", None,
            update_modified=True,
        )
        mock_commit.assert_called_once()

    def test_missing_user_id_is_rejected(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import delete_device_user

        with self.assertRaises(frappe.ValidationError):
            delete_device_user("Test-ZK-Device", "")


class TestDownloadDeviceUsersExcel(unittest.TestCase):
    """download_device_users_excel exports the dialog's fetched users as a workbook."""

    def _device(self, name="Test-ZK-Device", device_name="Test ZK Device"):
        device = MagicMock()
        device.name = name
        device.device_name = device_name
        return device

    def _users(self):
        return [
            {
                "user_id": "100", "name": "Abebe Kebede",
                "punch_methods": ["Fingerprint", "Password"],
                "employee": "HR-EMP-00001", "employee_name": "Abebe Kebede",
                "shift_type": "Morning",
            },
            {"user_id": "200", "name": "Unmapped User", "punch_methods": []},
        ]

    def _load(self, content):
        return load_workbook(BytesIO(base64.b64decode(content)))

    def test_returns_base64_workbook_matching_dialog_columns(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import download_device_users_excel

        with patch("frappe.get_doc", return_value=self._device()):
            result = download_device_users_excel(
                "Test-ZK-Device", users=json.dumps(self._users()))

        self.assertTrue(result["success"])
        self.assertEqual(result["filename"], "Device_Users_Test-ZK-Device.xlsx")

        ws = self._load(result["content"]).active
        self.assertEqual(ws.title, "Device Users")
        self.assertEqual(ws["A1"].value, "Employees On Test ZK Device")
        self.assertEqual(
            [ws.cell(row=4, column=c).value for c in range(1, 8)],
            ["Status", "Device ID", "Name On Device", "Punch Methods",
             "Mapped Employee", "Employee Name", "Shift Type"],
        )

        # Mapped user row
        self.assertEqual(ws.cell(row=5, column=1).value, "Mapped")
        self.assertEqual(ws.cell(row=5, column=2).value, "100")
        self.assertEqual(ws.cell(row=5, column=3).value, "Abebe Kebede")
        self.assertEqual(ws.cell(row=5, column=4).value, "Fingerprint, Password")
        self.assertEqual(ws.cell(row=5, column=5).value, "HR-EMP-00001")
        self.assertEqual(ws.cell(row=5, column=7).value, "Morning")

        # Unmapped user row
        self.assertEqual(ws.cell(row=6, column=1).value, "Not mapped")
        self.assertEqual(ws.cell(row=6, column=2).value, "200")
        # No enrolled methods: openpyxl stores the empty string as an empty cell
        self.assertIsNone(ws.cell(row=6, column=4).value)
        self.assertIsNone(ws.cell(row=7, column=1).value)

    def test_accepts_plain_list_payload(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import download_device_users_excel

        with patch("frappe.get_doc", return_value=self._device()):
            result = download_device_users_excel("Test-ZK-Device", users=self._users())

        self.assertTrue(result["success"])
        self.assertEqual(self._load(result["content"]).active["B5"].value, "100")

    def test_empty_users_produces_headers_only(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import download_device_users_excel

        with patch("frappe.get_doc", return_value=self._device()):
            result = download_device_users_excel("Test-ZK-Device", users=[])

        self.assertTrue(result["success"])
        ws = self._load(result["content"]).active
        self.assertEqual(ws["A4"].value, "Status")
        self.assertIsNone(ws["A5"].value)

    def test_rejects_invalid_json_payload(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import download_device_users_excel

        with patch("frappe.get_doc", return_value=self._device()), \
                self.assertRaises(frappe.ValidationError):
            download_device_users_excel("Test-ZK-Device", users="{bad json")

    def test_rejects_non_list_payload(self):
        from zkteco_attendance.zkteco_attendance.api.endpoints import download_device_users_excel

        with patch("frappe.get_doc", return_value=self._device()), \
                self.assertRaises(frappe.ValidationError):
            download_device_users_excel(
                "Test-ZK-Device", users=json.dumps({"user_id": "100"}))


if __name__ == "__main__":
    unittest.main()
