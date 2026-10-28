"""
Unit tests for the Daily Check-ins page Excel report builder.

The workbook builder runs server-side when the page's "Download Excel"
button is clicked. These tests keep the builder deterministic and safe:
two worksheets (employee daily data + summary), correct row content,
per-employee aggregates, the grand-total row, and the empty-data state.

Run with: bench run-tests --app zkteco_attendance
"""

import unittest

from openpyxl import load_workbook

from zkteco_attendance.zkteco_attendance.page.zk_daily_checkins.zk_daily_checkins import (
    _build_excel_workbook,
    _checkins_text,
    _employee_summary,
)


def _day(date_str, status="Present", hours=7.6, **overrides):
    day = {
        "date": date_str,
        "weekday": "Monday",
        "is_saturday": False,
        "is_weekend": False,
        "is_holiday": False,
        "holiday_name": "",
        "status": status,
        "hours": hours,
        "effective_hours": hours,
        "is_late": False,
        "is_early_exit": False,
        "late_minutes": 0.0,
        "early_minutes": 0.0,
        "overtime_hours": 0.0,
        "day_ot_hours": 0.0,
        "night_ot_hours": 0.0,
        "weekend_ot_hours": 0.0,
        "holiday_ot_hours": 0.0,
        "checkins": [],
    }
    day.update(overrides)
    return day


def _checkin(time_str, log_type="IN", **overrides):
    checkin = {
        "name": "CHK-1",
        "time": time_str,
        "log_type": log_type,
        "is_overtime": False,
        "manually_edited": False,
        "ignored": False,
    }
    checkin.update(overrides)
    return checkin


def _employee(employee="EMP-0001", employee_name="John Smith", days=None, **overrides):
    emp = {
        "employee": employee,
        "employee_name": employee_name,
        "department": "Production",
        "designation": "Operator",
        "zk_biometric_device": "DEV-01",
        "attendance_device_id": "101",
        "shift_type": "General",
        "days": days if days is not None else [],
    }
    emp.update(overrides)
    return emp


def _sample_data():
    return {
        "attendance_summary": "SUM-0001",
        "company": "Acme Inc",
        "from_date": "2026-08-01",
        "to_date": "2026-08-31",
        "employees": [
            _employee(days=[
                _day("2026-08-03", checkins=[_checkin("08:57:00", "IN")]),
                _day("2026-08-04", status="Half Day", hours=3.5,
                     is_late=True, late_minutes=12.0),
                _day("2026-08-09", status="Weekly Off", hours=0.0),
                _day("2026-08-14", status="Holiday", hours=0.0,
                     is_holiday=True, holiday_name="Independence Day"),
                _day("2026-08-20", hours=8.5, overtime_hours=0.5,
                     day_ot_hours=0.5,
                     checkins=[
                         _checkin("08:00:00", "IN"),
                         _checkin("17:30:00", "OUT"),
                     ]),
            ]),
            _employee(employee="EMP-0002", employee_name="Jane Roe",
                      department="Finance", designation="",
                      zk_biometric_device="DEV-02",
                      attendance_device_id="102", shift_type="",
                      days=[
                          _day("2026-08-03", hours=9.0, checkins=[
                              _checkin("08:00:00", "IN", manually_edited=True),
                              _checkin("09:15:00", "OUT", ignored=True),
                              _checkin("17:00:00", "OUT", is_overtime=True),
                          ]),
                      ]),
        ],
    }


def _load_workbook(data):
    return load_workbook(__import__("io").BytesIO(_build_excel_workbook(data)))


class TestCheckinsText(unittest.TestCase):

    def test_checkins_text_formats_punches(self):
        checkins = [
            _checkin("08:57:00", "IN"),
            _checkin("17:30:00", "OUT"),
        ]
        self.assertEqual(_checkins_text(checkins), "08:57 IN; 17:30 OUT")

    def test_checkins_text_marks_ot_manual_and_ignored(self):
        checkins = [
            _checkin("08:00:00", "IN", manually_edited=True),
            _checkin("09:15:00", "OUT", ignored=True),
            _checkin("17:00:00", "OUT", is_overtime=True),
        ]
        self.assertEqual(
            _checkins_text(checkins),
            "08:00 IN *; 09:15 OUT (ignored); 17:00 OUT (OT)",
        )

    def test_checkins_text_empty(self):
        self.assertEqual(_checkins_text([]), "")
        self.assertEqual(_checkins_text(None), "")


class TestEmployeeSummary(unittest.TestCase):

    def test_working_days_half_day_and_holiday(self):
        emp = _employee(days=[
            _day("2026-08-03"),                                   # Present  -> 1.0
            _day("2026-08-04", status="Half Day", hours=3.5),     # Half Day -> 0.5
            _day("2026-08-09", status="Weekly Off", hours=0.0),   # not counted
            _day("2026-08-14", status="Holiday", hours=0.0),      # Holiday -> 1.0
        ])
        s = _employee_summary(emp)
        self.assertEqual(s["working_days"], 2.5)
        self.assertEqual(s["absent_days"], 0)
        self.assertEqual(s["total_hours"], 11.1)

    def test_absent_invalid_review_counted(self):
        emp = _employee(days=[
            _day("2026-08-03", status="Absent", hours=0.0),
            _day("2026-08-04", status="Invalid", hours=0.0),
            _day("2026-08-05", status="Manual Review", hours=6.0),
        ])
        s = _employee_summary(emp)
        self.assertEqual(s["absent_days"], 1)
        self.assertEqual(s["invalid_days"], 1)
        self.assertEqual(s["review_days"], 1)
        self.assertEqual(s["working_days"], 0.0)
        self.assertEqual(s["total_hours"], 6.0)


class TestExcelWorkbook(unittest.TestCase):

    def test_workbook_has_two_sheets_with_headers(self):
        wb = _load_workbook(_sample_data())
        self.assertEqual(wb.sheetnames, ["Daily Checkins", "Summary"])

        daily = wb["Daily Checkins"]
        headers = [daily.cell(row=4, column=c).value for c in range(1, 33)]
        self.assertEqual(headers[0], "Employee ID")
        self.assertEqual(headers[7], "Date")
        self.assertEqual(headers[9], "Status")
        self.assertEqual(headers[10], "Lunch Dev (min)")
        self.assertEqual(headers[11], "Lunch Dev Flag")
        self.assertEqual(headers[12], "Add. Work (min)")
        self.assertEqual(headers[13], "Add: Before Start (min)")
        self.assertEqual(headers[16], "Add: After End (min)")
        self.assertEqual(headers[17], "Hours")
        self.assertEqual(headers[18], "Formatted Hours")
        self.assertEqual(headers[24], "Late (min)")
        self.assertEqual(headers[25], "Early (min)")
        self.assertEqual(headers[26], "Late: Shift Start (min)")
        self.assertEqual(headers[27], "Late: Lunch End (min)")
        self.assertEqual(headers[28], "Early: Lunch Start (min)")
        self.assertEqual(headers[29], "Early: Shift End (min)")
        self.assertEqual(headers[31], "Check-ins")

        summary = wb["Summary"]
        s_headers = [summary.cell(row=4, column=c).value for c in range(1, 27)]
        self.assertEqual(s_headers[0], "Employee ID")
        self.assertEqual(s_headers[7], "Working Days")
        self.assertEqual(s_headers[8], "Half Days")
        self.assertEqual(s_headers[9], "Absent Days")
        self.assertEqual(s_headers[10], "Late (min)")
        self.assertEqual(s_headers[11], "Early (min)")
        self.assertEqual(s_headers[12], "Late: Shift Start (min)")
        self.assertEqual(s_headers[13], "Late: Lunch End (min)")
        self.assertEqual(s_headers[14], "Lunch Dev Days")
        self.assertEqual(s_headers[15], "Lunch Dev (min)")
        self.assertEqual(s_headers[16], "Add. Work (min)")
        self.assertEqual(s_headers[17], "Add. Work (hrs)")
        self.assertEqual(s_headers[25], "Total OT")

    def test_daily_sheet_has_one_row_per_employee_day(self):
        wb = _load_workbook(_sample_data())
        daily = wb["Daily Checkins"]

        # 6 rows: EMP-0001 has 5 days, EMP-0002 has 1 day
        emp_ids = [daily.cell(row=r, column=1).value for r in range(5, 11)]
        self.assertEqual(emp_ids, ["EMP-0001"] * 5 + ["EMP-0002"])

        # First data row: EMP-0001, 2026-08-03, Present, 7.60 h
        self.assertEqual(daily.cell(row=5, column=2).value, "John Smith")
        self.assertEqual(daily.cell(row=5, column=8).value, "2026-08-03")
        self.assertEqual(daily.cell(row=5, column=10).value, "Present")
        self.assertAlmostEqual(daily.cell(row=5, column=18).value, 7.6, places=2)
        self.assertEqual(daily.cell(row=5, column=19).value, "7:36")
        self.assertEqual(daily.cell(row=5, column=32).value, "08:57 IN")

        # Half Day row carries the late minutes
        self.assertEqual(daily.cell(row=6, column=10).value, "Half Day")
        self.assertEqual(daily.cell(row=6, column=25).value, 12)

        # OT row: day OT in its own column + total OT
        self.assertAlmostEqual(daily.cell(row=9, column=20).value, 0.5, places=2)
        self.assertAlmostEqual(daily.cell(row=9, column=24).value, 0.5, places=2)

        # EMP-0002 row: check-in text marks manual / ignored / OT punches
        self.assertEqual(daily.cell(row=10, column=1).value, "EMP-0002")
        self.assertEqual(
            daily.cell(row=10, column=32).value,
            "08:00 IN *; 09:15 OUT (ignored); 17:00 OUT (OT)",
        )

    def test_daily_sheet_includes_lunch_deviation_and_additional_work(self):
        """Lunch deviation, its flag and the additional-working breakdown are
        exported per day, so the workbook matches the page columns."""
        wb = _load_workbook({
            "attendance_summary": "SUM-0001",
            "company": "Acme Inc",
            "from_date": "2026-08-03",
            "to_date": "2026-08-03",
            "employees": [_employee(days=[
                _day("2026-08-03", hours=8.5,
                     lunch_deviation=18.0, is_lunch_deviation=True,
                     additional_minutes=80.0,
                     additional_shift_start_minutes=10.0,
                     additional_lunch_start_minutes=5.0,
                     additional_lunch_end_minutes=0.0,
                     additional_shift_end_minutes=65.0,
                     is_late=True, late_minutes=17.0,
                     late_shift_minutes=12.0, late_lunch_minutes=5.0,
                     is_early_exit=True, early_minutes=9.0,
                     early_lunch_minutes=4.0, early_shift_minutes=5.0),
            ])],
        })
        daily = wb["Daily Checkins"]
        self.assertAlmostEqual(daily.cell(row=5, column=11).value, 18.0, places=1)
        self.assertEqual(daily.cell(row=5, column=12).value, "Yes")
        self.assertAlmostEqual(daily.cell(row=5, column=13).value, 80.0, places=1)
        self.assertAlmostEqual(daily.cell(row=5, column=14).value, 10.0, places=1)
        self.assertAlmostEqual(daily.cell(row=5, column=15).value, 5.0, places=1)
        self.assertAlmostEqual(daily.cell(row=5, column=16).value, 0.0, places=1)
        self.assertAlmostEqual(daily.cell(row=5, column=17).value, 65.0, places=1)
        self.assertEqual(daily.cell(row=5, column=25).value, 17)
        self.assertEqual(daily.cell(row=5, column=26).value, 9)
        self.assertAlmostEqual(daily.cell(row=5, column=27).value, 12.0, places=1)
        self.assertAlmostEqual(daily.cell(row=5, column=28).value, 5.0, places=1)
        self.assertAlmostEqual(daily.cell(row=5, column=29).value, 4.0, places=1)
        self.assertAlmostEqual(daily.cell(row=5, column=30).value, 5.0, places=1)

        # The summary sheet carries the period totals of the same figures.
        summary = wb["Summary"]
        self.assertEqual(summary.cell(row=5, column=15).value, 1)      # Lunch Dev Days
        self.assertAlmostEqual(summary.cell(row=5, column=16).value, 18.0, places=1)
        self.assertAlmostEqual(summary.cell(row=5, column=17).value, 80.0, places=1)
        self.assertAlmostEqual(summary.cell(row=5, column=18).value, 1.33, places=2)
        self.assertEqual(summary.cell(row=5, column=9).value, 0)       # Half Days

    def test_summary_sheet_aggregates_per_employee(self):
        wb = _load_workbook(_sample_data())
        summary = wb["Summary"]

        # Row 5 = EMP-0001: Present(03) + Half Day(04) + Holiday(14) + Present(20)
        #                = 1.0 + 0.5 + 1.0 + 1.0 = 3.5 days
        self.assertEqual(summary.cell(row=5, column=1).value, "EMP-0001")
        self.assertAlmostEqual(summary.cell(row=5, column=8).value, 3.5, places=2)
        self.assertEqual(summary.cell(row=5, column=9).value, 1)  # Half Days
        # Late/early minutes sit after Half / Absent Days
        self.assertEqual(summary.cell(row=5, column=11).value, 12)  # Half Day late
        self.assertEqual(summary.cell(row=5, column=12).value, 0)
        self.assertAlmostEqual(summary.cell(row=5, column=21).value, 19.6, places=2)
        self.assertAlmostEqual(summary.cell(row=5, column=22).value, 0.5, places=2)
        self.assertAlmostEqual(summary.cell(row=5, column=26).value, 0.5, places=2)

        # Row 6 = EMP-0002: 1 working day, 9.0 h
        self.assertEqual(summary.cell(row=6, column=1).value, "EMP-0002")
        self.assertAlmostEqual(summary.cell(row=6, column=8).value, 1.0, places=2)
        self.assertEqual(summary.cell(row=6, column=11).value, 0)
        self.assertAlmostEqual(summary.cell(row=6, column=21).value, 9.0, places=2)

        # Row 7 = grand totals
        self.assertEqual(summary.cell(row=7, column=1).value, "Grand Total")
        self.assertAlmostEqual(summary.cell(row=7, column=8).value, 4.5, places=2)
        self.assertEqual(summary.cell(row=7, column=11).value, 12)
        self.assertEqual(summary.cell(row=7, column=12).value, 0)
        self.assertAlmostEqual(summary.cell(row=7, column=21).value, 28.6, places=2)
        self.assertAlmostEqual(summary.cell(row=7, column=26).value, 0.5, places=2)

    def test_empty_state_produces_headers_only(self):
        wb = _load_workbook({
            "attendance_summary": "",
            "company": "",
            "from_date": "2026-08-01",
            "to_date": "2026-08-31",
            "employees": [],
        })
        self.assertEqual(wb.sheetnames, ["Daily Checkins", "Summary"])
        daily = wb["Daily Checkins"]
        self.assertEqual(daily.cell(row=4, column=1).value, "Employee ID")
        # No data rows after the header
        self.assertIsNone(daily.cell(row=5, column=1).value)
        summary = wb["Summary"]
        self.assertIsNone(summary.cell(row=5, column=1).value)
        self.assertNotEqual(summary.cell(row=5, column=1).value, "Grand Total")


if __name__ == "__main__":
    unittest.main()