"""
Unit tests for the Daily Check-ins "Check lunch deviation" dialog.

Covers the payload that lists employees whose lunch deviation exceeded their
shift's threshold (with the counted days in one comma separated column) and the
Excel workbook that mirrors the dialog.

Run with: bench run-tests --app zkteco_attendance
"""

import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from zkteco_attendance.zkteco_attendance.page.zk_daily_checkins.zk_daily_checkins import (
    _build_lunch_deviation_workbook,
    get_lunch_deviation_flags,
)

ENDPOINTS = "zkteco_attendance.zkteco_attendance.api.endpoints.get_daily_checkins"


def _day(date_str, deviation=0.0, flagged=False):
    return {
        "date": date_str,
        "status": "Present",
        "lunch_deviation": deviation,
        "is_lunch_deviation": flagged,
    }


def _employee(employee, days, threshold=5, enabled=True, **overrides):
    """Employee row as `get_daily_checkins` builds it — the flag follows the threshold."""
    deviated = [d for d in days if d.get("is_lunch_deviation")]
    count = len(deviated) if enabled else 0
    emp = {
        "employee": employee,
        "employee_name": employee.upper(),
        "department": "Production",
        "designation": "Operator",
        "shift_type": "General",
        "lunch_deviation_enabled": enabled,
        "lunch_deviation_days": count,
        "lunch_deviation_minutes": round(
            sum(d.get("lunch_deviation") or 0 for d in deviated), 1) if enabled else 0.0,
        "lunch_deviation_threshold": threshold,
        "lunch_deviation_flagged": enabled and count > threshold,
        "days": days,
    }
    emp.update(overrides)
    return emp


def _payload(employees):
    return {
        "from_date": "2026-08-01",
        "to_date": "2026-08-31",
        "employees": employees,
    }


class TestLunchDeviationFlags(unittest.TestCase):
    """The dialog payload only lists employees above their threshold."""

    @patch(ENDPOINTS)
    def test_only_flagged_employees_are_returned(self, mock_checkins):
        mock_checkins.return_value = _payload([
            _employee("EMP-A", [_day("2026-08-03", 25.0, True),
                                _day("2026-08-04", 0.0, False)], threshold=0),
            _employee("EMP-B", [_day("2026-08-03", 15.0, True)], threshold=5),
            _employee("EMP-C", [_day("2026-08-03", 0.0, False)], enabled=False),
            _employee("EMP-D", [_day("2026-08-03", 40.0, True),
                                _day("2026-08-04", 20.0, True)], threshold=1),
        ])

        result = get_lunch_deviation_flags(from_date="2026-08-01", to_date="2026-08-31")

        self.assertEqual([e["employee"] for e in result["flagged"]], ["EMP-A", "EMP-D"])
        self.assertEqual(result["total_employees_checked"], 4)
        self.assertEqual(result["from_date"], "2026-08-01")
        self.assertEqual(result["to_date"], "2026-08-31")

    @patch(ENDPOINTS)
    def test_deviated_days_are_comma_separated_in_one_value(self, mock_checkins):
        mock_checkins.return_value = _payload([
            _employee("EMP-A", [
                _day("2026-08-03", 25.0, True),
                _day("2026-08-04", 0.0, False),
                _day("2026-08-05", 12.0, True),
                _day("2026-08-06", 18.0, True),
                _day("2026-08-07", 5.0, True),
                _day("2026-08-10", 9.0, True),
                _day("2026-08-11", 11.0, True),
            ], threshold=5),
        ])

        emp = get_lunch_deviation_flags()["flagged"][0]

        self.assertEqual(emp["deviation_days"], 6)
        self.assertEqual(emp["threshold"], 5)
        self.assertAlmostEqual(emp["deviation_minutes"], 80.0, places=1)
        self.assertEqual(emp["deviation_dates_text"],
                         "2026-08-03, 2026-08-05, 2026-08-06, 2026-08-07, "
                         "2026-08-10, 2026-08-11")

    @patch(ENDPOINTS)
    def test_no_flags_returns_an_empty_list(self, mock_checkins):
        mock_checkins.return_value = _payload([
            _employee("EMP-B", [_day("2026-08-03", 15.0, True)], threshold=5),
        ])

        result = get_lunch_deviation_flags()

        self.assertEqual(result["flagged"], [])
        self.assertEqual(result["total_employees_checked"], 1)


class TestLunchDeviationWorkbook(unittest.TestCase):
    """The Excel download mirrors the dialog: one row per flagged employee."""

    def _workbook(self):
        return load_workbook(BytesIO(_build_lunch_deviation_workbook({
            "from_date": "2026-08-01",
            "to_date": "2026-08-31",
            "flagged": [
                {
                    "employee": "EMP-A", "employee_name": "Asha",
                    "department": "Production", "shift_type": "General",
                    "threshold": 5, "deviation_days": 6,
                    "deviation_minutes": 120.5,
                    "deviation_dates": ["2026-08-03", "2026-08-05",
                                        "2026-08-06", "2026-08-07",
                                        "2026-08-10", "2026-08-11"],
                    "deviation_dates_text": "2026-08-03, 2026-08-05, 2026-08-06, "
                                            "2026-08-07, 2026-08-10, 2026-08-11",
                },
                {
                    "employee": "EMP-C", "employee_name": "Chen",
                    "department": "", "shift_type": "Night",
                    "threshold": 1, "deviation_days": 2,
                    "deviation_minutes": 61.0,
                    "deviation_dates": ["2026-08-03", "2026-08-04"],
                    "deviation_dates_text": "2026-08-03, 2026-08-04",
                },
            ],
        }))).active

    def test_header_row_labels_the_columns(self):
        ws = self._workbook()

        self.assertIn("Lunch Deviation", ws.title)
        self.assertEqual(
            [ws.cell(row=4, column=i).value for i in range(1, 9)],
            ["Employee ID", "Employee Name", "Department", "Shift",
             "Deviation Days", "Threshold (Days)", "Total Deviation (min)",
             "Deviation Dates"],
        )

    def test_dates_are_comma_separated_in_a_single_cell(self):
        ws = self._workbook()

        self.assertEqual(ws.max_row, 6)
        self.assertEqual(ws.cell(row=5, column=2).value, "Asha")
        self.assertEqual(ws.cell(row=5, column=5).value, 6)
        self.assertEqual(ws.cell(row=5, column=6).value, 5)
        self.assertAlmostEqual(ws.cell(row=5, column=7).value, 120.5, places=1)
        self.assertEqual(
            ws.cell(row=5, column=8).value,
            "2026-08-03, 2026-08-05, 2026-08-06, 2026-08-07, 2026-08-10, 2026-08-11",
        )
        self.assertEqual(ws.cell(row=6, column=8).value, "2026-08-03, 2026-08-04")

    def test_empty_payload_writes_headers_only(self):
        ws = load_workbook(BytesIO(_build_lunch_deviation_workbook({
            "from_date": "2026-08-01",
            "to_date": "2026-08-31",
            "flagged": [],
        }))).active

        self.assertEqual(ws.max_row, 4)


if __name__ == "__main__":
    unittest.main()
