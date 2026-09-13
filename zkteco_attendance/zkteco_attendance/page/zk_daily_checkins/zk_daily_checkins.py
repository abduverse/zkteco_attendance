import frappe
from frappe import _

from frappe.utils import cstr, flt, getdate


@frappe.whitelist()
def get_data(attendance_summary=None, from_date=None, to_date=None,
             employee_list=None, company=None, biometric_device=None,
             filter_employee=None):
    from zkteco_attendance.zkteco_attendance.api.endpoints import get_daily_checkins
    return get_daily_checkins(
        attendance_summary=attendance_summary,
        from_date=from_date,
        to_date=to_date,
        employee_list=employee_list,
        company=company,
        biometric_device=biometric_device,
        filter_employee=filter_employee,
    )


@frappe.whitelist()
def save_manual_checkin(attendance_summary=None, employee=None, checkin_time=None,
                        log_type=None, checkin_name=None, is_overtime=0, remark=None):
    from zkteco_attendance.zkteco_attendance.api.endpoints import save_manual_checkin as _save
    return _save(attendance_summary=attendance_summary, employee=employee,
                 checkin_time=checkin_time, log_type=log_type,
                 checkin_name=checkin_name, is_overtime=is_overtime, remark=remark)


@frappe.whitelist()
def create_manual_checkin_request(employee=None, checkin_date=None, checkin_time=None,
                                  log_type=None, is_overtime=0, attendance_summary=None,
                                  checkin_name=None, remarks=None):
    from zkteco_attendance.zkteco_attendance.api.endpoints import create_manual_checkin_request as _create
    return _create(employee=employee, checkin_date=checkin_date, checkin_time=checkin_time,
                   log_type=log_type, is_overtime=is_overtime, attendance_summary=attendance_summary,
                   checkin_name=checkin_name, remarks=remarks)


@frappe.whitelist()
def toggle_ignore_checkin(checkin_name):
    from zkteco_attendance.zkteco_attendance.api.endpoints import toggle_ignore_checkin as _toggle
    return _toggle(checkin_name=checkin_name)


@frappe.whitelist()
def get_invalid_days(attendance_summary=None, from_date=None, to_date=None,
                     employee_list=None, company=None, biometric_device=None,
                     filter_employee=None):
    """
    Return only the employees whose daily breakdown contains "Invalid" days,
    together with the exact invalid dates and counts.

    Accepts the same arguments as `get_data`, so "Check invalids" always
    matches the currently filtered view on the page.
    """
    from zkteco_attendance.zkteco_attendance.api.endpoints import get_daily_checkins

    data = get_daily_checkins(
        attendance_summary=attendance_summary,
        from_date=from_date,
        to_date=to_date,
        employee_list=employee_list,
        company=company,
        biometric_device=biometric_device,
        filter_employee=filter_employee,
    )

    invalids = []
    for emp in data.get("employees") or []:
        invalid_dates = [
            d.get("date")
            for d in (emp.get("days") or [])
            if (d.get("status") or "") == "Invalid"
        ]
        if not invalid_dates:
            continue

        # Include the raw punches recorded on each invalid day so the dialog
        # can show exactly which checkins are missing their IN/OUT pair.
        day_map = {d.get("date"): d for d in (emp.get("days") or [])}
        day_checkins = [
            {
                "date":     dt,
                "checkins": (day_map.get(dt) or {}).get("checkins") or [],
            }
            for dt in invalid_dates
        ]

        invalids.append({
            "employee":      emp.get("employee"),
            "employee_name": emp.get("fullname") or emp.get("employee_name") or emp.get("employee"),
            "department":    emp.get("department") or "",
            "shift_type":    emp.get("shift_type") or "",
            "invalid_count": len(invalid_dates),
            "invalid_dates": invalid_dates,
            "day_checkins":  day_checkins,
        })

    return {
        "from_date": data.get("from_date"),
        "to_date":   data.get("to_date"),
        "invalids":  invalids,
        "total_employees_checked": len(data.get("employees") or []),
    }


@frappe.whitelist()
def get_employee_shift_info(employee, work_date=None):
    from zkteco_attendance.zkteco_attendance.api.endpoints import get_employee_shift_info
    return get_employee_shift_info(employee, work_date)


# ---------------------------------------------------------------------------
# PDF download — renders the loaded Daily Checkins view as a styled PDF
# ---------------------------------------------------------------------------

@frappe.whitelist()
def download_pdf(attendance_summary=None, from_date=None, to_date=None,
                 employee_list=None, company=None, biometric_device=None,
                 filter_employee=None):
    """
    Render the currently filtered Daily Checkins view as a styled PDF and
    stream it to the browser as a file download.

    Accepts the exact same arguments as `get_data`, so the PDF always
    matches what the page's Load button displays.
    """
    from zkteco_attendance.zkteco_attendance.api.endpoints import get_daily_checkins
    from frappe.utils.pdf import get_pdf

    data = get_daily_checkins(
        attendance_summary=attendance_summary,
        from_date=from_date,
        to_date=to_date,
        employee_list=employee_list,
        company=company,
        biometric_device=biometric_device,
        filter_employee=filter_employee,
    )

    pdf = get_pdf(_render_pdf_html(data))

    from_d = cstr(data.get("from_date") or from_date or "start")
    to_d   = cstr(data.get("to_date")   or to_date   or "end")

    frappe.local.response.type        = "pdf"
    frappe.local.response.filename    = "Daily_Checkins_{0}_to_{1}.pdf".format(from_d, to_d)
    frappe.local.response.filecontent = pdf


# ---------------------------------------------------------------------------
# Excel download — streams the loaded Daily Checkins view as an .xlsx workbook
# Sheet 1 "Daily Checkins": one row per employee per day (the employee tables).
# Sheet 2 "Summary": per-employee aggregates + grand totals.
# ---------------------------------------------------------------------------

@frappe.whitelist()
def download_excel(attendance_summary=None, from_date=None, to_date=None,
                   employee_list=None, company=None, biometric_device=None,
                   filter_employee=None):
    """
    Stream the currently filtered Daily Checkins view as an Excel workbook.
    Accepts the exact same arguments as `get_data`, so the workbook always
    matches what the page's Load button displays.
    """
    from zkteco_attendance.zkteco_attendance.api.endpoints import get_daily_checkins

    data = get_daily_checkins(
        attendance_summary=attendance_summary,
        from_date=from_date,
        to_date=to_date,
        employee_list=employee_list,
        company=company,
        biometric_device=biometric_device,
        filter_employee=filter_employee,
    )

    xlsx = _build_excel_workbook(data)

    from_d = cstr(data.get("from_date") or from_date or "start")
    to_d   = cstr(data.get("to_date")   or to_date   or "end")

    frappe.local.response.type        = "binary"
    frappe.local.response.filename    = "Daily_Checkins_{0}_to_{1}.xlsx".format(from_d, to_d)
    frappe.local.response.filecontent = xlsx


# ── HTML/CSS for the PDF document ────────────────────────────────────────────

_PDF_CSS = """
* { margin: 0; padding: 0; }
body {
    font-family: "Segoe UI", Roboto, "Helvetica Neue", Helvetica, Arial, sans-serif;
    color: #26323c;
    font-size: 8pt;
    background: #ffffff;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
}
table { border-collapse: collapse; width: 100%; }

/* ── Document header ─────────────────────────────────────────────────────── */
.doc-head {
    background: #16324f;
    background: linear-gradient(120deg, #12293f 0%, #1d4770 55%, #2a6699 100%);
    border-radius: 8px 8px 0 0;
    color: #ffffff;
    padding: 11px 14px 12px;
}
.doc-title { font-size: 15pt; font-weight: 700; letter-spacing: 0.2px; }
.doc-sub { font-size: 7.5pt; opacity: 0.85; margin-top: 2px; }
.doc-head td { padding: 0; }
.head-right { text-align: right; }
.period-box {
    display: inline-block;
    background: rgba(255, 255, 255, 0.14);
    border: 1px solid rgba(255, 255, 255, 0.45);
    border-radius: 7px;
    padding: 5px 12px;
    text-align: right;
}
.period-box .lbl { font-size: 6pt; text-transform: uppercase; letter-spacing: 1.1px; opacity: 0.8; }
.period-box .val { font-size: 10pt; font-weight: 700; margin-top: 1px; }

.meta-strip {
    background: #eef4fb;
    border: 1px solid #d6e2ef;
    border-top: 0;
    padding: 7px 10px;
}
.meta-chip {
    display: inline-block;
    background: #ffffff;
    border: 1px solid #cfdbe8;
    border-radius: 11px;
    padding: 1px 8px;
    margin: 1px 5px 1px 0;
    font-size: 7pt;
    color: #3f5264;
}
.meta-chip b { color: #16324f; }

.legend {
    background: #f8fafc;
    border: 1px solid #e0e7ef;
    border-top: 0;
    padding: 6px 10px 7px;
    font-size: 6.6pt;
    color: #5b6b7a;
}
.legend td { padding: 2px 14px 2px 0; vertical-align: top; }
.legend .lg-label {
    text-transform: uppercase;
    letter-spacing: 0.9px;
    font-weight: 700;
    color: #16324f;
    font-size: 5.9pt;
    margin-bottom: 3px;
}
.legend .lg-sub { margin-top: 2px; }

/* ── Status pills ────────────────────────────────────────────────────────── */
.pill {
    display: inline-block;
    border-radius: 8px;
    padding: 1px 7px;
    font-size: 6.8pt;
    font-weight: 700;
    border: 1px solid;
    white-space: nowrap;
}
.p-present { background: #e4f6ea; border-color: #a9d8bc; color: #1e7a3e; }
.p-half    { background: #fff3da; border-color: #f0d396; color: #9a6b06; }
.p-absent  { background: #fdeaea; border-color: #f0b8b8; color: #b3261e; }
.p-invalid { background: #eceff1; border-color: #c8ced4; color: #5c6a75; }
.p-review  { background: #e8f0fd; border-color: #b7cdea; color: #1c56a3; }
.p-holiday { background: #f3ecfb; border-color: #d6c2ea; color: #6b36ae; }
.p-weekoff { background: #f1f1f2; border-color: #d6d6d8; color: #6e6e70; }

.warn { font-size: 6.4pt; font-weight: 700; color: #a05a12; margin-left: 4px; white-space: nowrap; }
.subnote { font-size: 6.4pt; color: #8a5a1a; margin-top: 1px; }
.muted { color: #9aa5ae; }

/* ── Overtime chips (D / N / W / H) ──────────────────────────────────────── */
.ot {
    display: inline-block;
    padding: 1px 5px;
    border-radius: 3px;
    font-weight: 700;
    font-size: 6.6pt;
    margin: 1px 1px 1px 0;
    border: 1px solid;
    white-space: nowrap;
}
.ot-day     { background: #e6eefc; border-color: #a9c6ea; color: #1c56a3; }
.ot-night   { background: #f2eafa; border-color: #cdb5e6; color: #6b36ae; }
.ot-weekend { background: #fdefdd; border-color: #e0b98a; color: #a05a12; }
.ot-holiday { background: #fbe7e7; border-color: #e5abab; color: #b3261e; }

/* ── Check-in chips ──────────────────────────────────────────────────────── */
.ck {
    display: inline-block;
    padding: 1px 5px;
    border-radius: 8px;
    border: 1px solid;
    font-size: 6.8pt;
    margin: 1px 2px 1px 0;
    white-space: nowrap;
}
.ck b { font-weight: 700; }
.ck-in  { background: #e6f6ec; border-color: #aed9bd; color: #1c7a3e; }
.ck-out { background: #fdeaea; border-color: #eeb9b9; color: #b3261e; }
.ck-ot  { background: #fff0e0; border-color: #e2b57f; color: #9a5700; }
.ck-ign { background: #f3f4f5; color: #77818a; border-style: dashed; }
.ck-mark { color: #1c56a3; font-weight: 800; }

/* ── Employee blocks ─────────────────────────────────────────────────────── */
.emp-block { margin-top: 14px; }
.emp-head {
    background: #16324f;
    background: linear-gradient(120deg, #16324f 0%, #1d4770 60%, #235a8c 100%);
    color: #ffffff;
    border-radius: 6px 6px 0 0;
    padding: 8px 12px;
}
.emp-head td { padding: 0; }
.emp-name { font-size: 11pt; font-weight: 700; }
.emp-meta { font-size: 7pt; opacity: 0.8; margin-top: 2px; }
.emp-chip-cell { text-align: right; }
.emp-badge {
    display: inline-block;
    border: 1px solid rgba(255, 255, 255, 0.45);
    border-radius: 10px;
    padding: 1px 8px;
    font-size: 6.8pt;
    margin: 1px 0 1px 6px;
    background: rgba(255, 255, 255, 0.12);
    color: #e9f1f9;
    white-space: nowrap;
}
.emp-badge b { color: #ffffff; }

.daily { margin: 0; }
.daily th {
    background: #1e4a73;
    color: #ffffff;
    font-size: 6.5pt;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    padding: 4px 6px;
    text-align: left;
    border: 1px solid #16324f;
}
.daily td {
    padding: 3px 6px;
    border: 1px solid #dfe6ee;
    vertical-align: middle;
}
.daily tbody tr:nth-child(even) td { background: #f5f8fc; }
.daily tbody tr { page-break-inside: avoid; }
.daily .num { text-align: right; white-space: nowrap; }
.daily .c-date { white-space: nowrap; }
.daily .ot-cell, .daily .ck-cell { line-height: 1.55; }
.daily tr.totals td {
    background: #e6edf5 !important;
    border-top: 2px solid #b6c4d4;
    font-weight: 700;
    color: #16324f;
    font-size: 7.2pt;
}

/* ── Empty state ─────────────────────────────────────────────────────────── */
.no-data {
    margin-top: 18px;
    border: 1px dashed #c3ced9;
    border-radius: 8px;
    background: #f8fafc;
    padding: 46px 20px;
    text-align: center;
    color: #7c8a98;
    font-size: 9pt;
}
"""


# ── Small rendering helpers ─────────────────────────────────────────────────

def _esc(value):
    """HTML-escape any value before it is embedded in the report."""
    return frappe.utils.escape_html(cstr(value or ""))


def _dec(hours, digits=2):
    """Decimal hours as a fixed-width string, e.g. 7.6 -> '7.60'."""
    return "{0:.{1}f}".format(flt(hours or 0), digits)


def _hhmm(hours):
    """Decimal hours to clock-style H:MM, e.g. 7.6 -> '7:36'.
    Mirrors the Formatted Hours column on the page."""
    total_minutes = max(0, int(round(flt(hours or 0) * 60)))
    h, m = divmod(total_minutes, 60)
    return "{0}:{1:02d}".format(h, m)


def _fmt_date(value):
    """'2026-08-01' -> '01 Aug 2026' (unchanged if it cannot be parsed)."""
    try:
        return getdate(value).strftime("%d %b %Y")
    except Exception:
        return _esc(value)


# Category metadata: (day key, letter, css class, legend label)
_OT_META = (
    ("day_ot_hours",     "D", "ot-day",     _("Day OT")),
    ("night_ot_hours",   "N", "ot-night",   _("Night OT")),
    ("weekend_ot_hours", "W", "ot-weekend", _("Weekend OT")),
    ("holiday_ot_hours", "H", "ot-holiday", _("Holiday OT")),
)

_STATUS_CLASS = {
    "Present":       "p-present",
    "Half Day":      "p-half",
    "Absent":        "p-absent",
    "Invalid":       "p-invalid",
    "Manual Review": "p-review",
    "Holiday":       "p-holiday",
    "Weekly Off":    "p-weekoff",
}


def _status_cell(day):
    status = cstr(day.get("status") or "")
    cls    = _STATUS_CLASS.get(status, "p-weekoff")
    parts  = ['<span class="pill {0}">{1}</span>'.format(cls, _esc(_(status) if status else ""))]

    marks = []
    if day.get("is_late"):
        marks.append(_("L-EN") + " " + str(int(flt(day.get("late_minutes") or 0))) + "m")
    if day.get("is_early_exit"):
        marks.append(_("E-EX") + " " + str(int(flt(day.get("early_minutes") or 0))) + "m")
    if marks:
        parts.append('<span class="warn">{0}</span>'.format(_esc(" ".join(marks))))

    if day.get("is_holiday") and day.get("holiday_name"):
        parts.append('<div class="subnote">{0}</div>'.format(_esc(day.get("holiday_name"))))
    return "".join(parts)


def _ot_html(row):
    """OT chips for a day dict or an aggregated totals dict."""
    chips = []
    for key, letter, cls, _title in _OT_META:
        value = flt(row.get(key) or 0)
        if value > 0:
            chips.append('<span class="ot {0}">{1} {2}</span>'.format(cls, letter, _dec(value, 1)))
    return "".join(chips) if chips else '<span class="muted">&mdash;</span>'


def _checkin_chips(day):
    checkins = day.get("checkins") or []
    if not checkins:
        return '<span class="muted">&mdash;</span>'

    parts = []
    for c in checkins:
        log_type = cstr(c.get("log_type") or "IN").upper()
        is_ot    = bool(c.get("is_overtime"))
        if is_ot:
            kind = "ck-ot"
        elif log_type == "IN":
            kind = "ck-in"
        else:
            kind = "ck-out"

        time_txt = cstr(c.get("time") or "")[:5]
        label    = log_type + (" " + _("OT") if is_ot else "")
        body     = '{0} <b>{1}</b>'.format(_esc(time_txt), _esc(label))
        if c.get("manually_edited"):
            body += ' <span class="ck-mark">*</span>'
        if c.get("ignored"):
            body = '<s>{0}</s>'.format(body)
            kind += " ck-ign"
        parts.append('<span class="ck {0}">{1}</span>'.format(kind, body))
    return " ".join(parts)


# ── Report building ──────────────────────────────────────────────────────────

def _employee_totals(days):
    totals = {"hours": 0.0}
    for key, _letter, _cls, _title in _OT_META:
        totals[key] = 0.0
    for d in days:
        totals["hours"] += flt(d.get("hours") or 0)
        for key, _letter, _cls, _title in _OT_META:
            totals[key] += flt(d.get(key) or 0)
    return totals


def _employee_table(emp):
    days = emp.get("days") or []
    rows = []
    for d in days:
        rows.append(
            '<tr>'
            '<td class="c-date">{0}</td>'
            '<td>{1}</td>'
            '<td>{2}</td>'
            '<td class="num">{3}</td>'
            '<td class="num">{4}</td>'
            '<td class="ot-cell">{5}</td>'
            '<td class="ck-cell">{6}</td>'
            '</tr>'.format(
                _fmt_date(d.get("date")),
                _esc(_(cstr(d.get("weekday") or ""))),
                _status_cell(d),
                _dec(d.get("hours")),
                _hhmm(d.get("hours")),
                _ot_html(d),
                _checkin_chips(d),
            )
        )

    totals_html = ""
    if days:
        totals = _employee_totals(days)
        totals_html = (
            '<tr class="totals">'
            '<td colspan="3">{0}</td>'
            '<td class="num">{1}</td>'
            '<td class="num">{2}</td>'
            '<td class="ot-cell">{3}</td>'
            '<td></td>'
            '</tr>'.format(
                _("Period total"),
                _dec(totals["hours"]),
                _hhmm(totals["hours"]),
                _ot_html(totals),
            )
        )

    return (
        '<table class="daily">'
        '<thead><tr>'
        '<th>{0}</th><th>{1}</th><th>{2}</th>'
        '<th class="num">{3}</th><th class="num">{4}</th>'
        '<th>{5}</th><th>{6}</th>'
        '</tr></thead>'
        '<tbody>{7}{8}</tbody>'
        '</table>'.format(
            _("Date"), _("Day"), _("Status"),
            _("Hours"), _("Formatted Hours"),
            _("OT Breakdown"), _("Check-ins"),
            "".join(rows), totals_html,
        )
    )


def _employee_block(emp):
    name = (emp.get("fullname") or emp.get("employee_name") or emp.get("employee") or "")

    meta_parts = []
    if emp.get("employee"):
        meta_parts.append(_esc(emp.get("employee")))
    if emp.get("department"):
        meta_parts.append(_esc(emp.get("department")))
    if emp.get("designation"):
        meta_parts.append(_esc(emp.get("designation")))

    badges = []
    if emp.get("zk_biometric_device"):
        badges.append('<span class="emp-badge"><b>{0}</b> {1}</span>'.format(
            _("Device"), _esc(emp.get("zk_biometric_device"))))
    if emp.get("attendance_device_id"):
        badges.append('<span class="emp-badge"><b>{0}</b> {1}</span>'.format(
            _("Device ID"), _esc(emp.get("attendance_device_id"))))
    if emp.get("shift_type"):
        badges.append('<span class="emp-badge"><b>{0}</b> {1}</span>'.format(
            _("Shift"), _esc(emp.get("shift_type"))))

    return (
        '<div class="emp-block">'
        '<div class="emp-head"><table class="layout"><tr>'
        '<td class="emp-title-cell">'
        '<div class="emp-name">{0}</div>'
        '<div class="emp-meta">{1}</div>'
        '</td>'
        '<td class="emp-chip-cell">{2}</td>'
        '</tr></table></div>'
        '{3}'
        '</div>'.format(_esc(name), " &middot; ".join(meta_parts), "".join(badges),
                        _employee_table(emp))
    )


def _render_pdf_html(data):
    """Build the full HTML document that gets converted to PDF."""
    from frappe.utils import now_datetime

    employees   = data.get("employees") or []
    from_date   = _fmt_date(data.get("from_date"))
    to_date     = _fmt_date(data.get("to_date"))
    period      = (from_date + " &ndash; " + to_date) if from_date != to_date else from_date

    # ── Meta chips ─────────────────────────────────────────────────────────
    def chip(label, value):
        return '<span class="meta-chip"><b>{0}:</b> {1}</span>'.format(_esc(label), _esc(value))

    meta = []
    if data.get("company"):
        meta.append(chip(_("Company"), data.get("company")))
    if data.get("attendance_summary"):
        meta.append(chip(_("Attendance Summary"), data.get("attendance_summary")))

    grand_hours = 0.0
    grand_ot    = 0.0
    for emp in employees:
        for d in emp.get("days") or []:
            grand_hours += flt(d.get("hours") or 0)
            grand_ot    += flt(d.get("overtime_hours") or 0)

    meta.append(chip(_("Employees"), str(len(employees))))
    if employees:
        meta.append(chip(_("Total Hours"), _dec(grand_hours) + " h (" + _hhmm(grand_hours) + ")"))
        if grand_ot > 0:
            meta.append(chip(_("Total OT"), _dec(grand_ot, 1) + " h"))

    generated    = frappe.session.user
    generated_at = now_datetime().strftime("%d %b %Y, %H:%M")
    meta.append(chip(_("Generated by"), generated + " \u00b7 " + generated_at))

    # ── Legend ─────────────────────────────────────────────────────────────
    status_samples = " ".join(
        '<span class="pill {0}">{1}</span>'.format(cls, _esc(_(label)))
        for label, cls in (
            ("Present", "p-present"), ("Half Day", "p-half"), ("Absent", "p-absent"),
            ("Manual Review", "p-review"), ("Holiday", "p-holiday"), ("Weekly Off", "p-weekoff"),
        )
    )

    legend = (
        '<table class="legend-table"><tr>'
        '<td width="33%"><div class="lg-label">{0}</div>{1}</td>'
        '<td width="33%"><div class="lg-label">{2}</div>{3}<div class="lg-sub">{4}</div></td>'
        '<td width="34%"><div class="lg-label">{5}</div>{6}</td>'
        '</tr></table>'.format(
            _("Status"), status_samples,
            _("Columns"),
            _("Hours = decimal hours, Formatted Hours = clock time (e.g. 7.60 / 7:36)."),
            _("OT: D = Day, N = Night, W = Weekend, H = Holiday overtime."),
            _("Markers"),
            _("L-EN / E-EX = late entry / early exit beyond grace. <b>*</b> = manually edited. <s>struck out</s> = ignored check-in."),
        )
    )

    # ── Body ───────────────────────────────────────────────────────────────
    if employees:
        blocks = "".join(_employee_block(emp) for emp in employees)
    else:
        blocks = '<div class="no-data">{0}</div>'.format(
            _("No attendance data found for the selected filters."))

    return (
        '<!DOCTYPE html><html><head><meta charset="utf-8">'
        '<style>{0}</style>'
        '</head><body><div class="page">'
        '<div class="doc-head"><table class="layout"><tr>'
        '<td><div class="doc-title">{1}</div>'
        '<div class="doc-sub">{2}</div></td>'
        '<td class="head-right"><div class="period-box">'
        '<div class="lbl">{3}</div><div class="val">{4}</div>'
        '</div></td></tr></table></div>'
        '<div class="meta-strip">{5}</div>'
        '<div class="legend">{6}</div>'
        '{7}'
        '</div></body></html>'.format(
            _PDF_CSS,
            _esc(_("Employee Daily Check-ins")),
            _esc(_("Biometric Attendance Report")),
            _esc(_("Report Period")),
            _esc(period),
            "".join(meta),
            legend,
            blocks,
        )
    )


# ── Excel workbook builder ──────────────────────────────────────────────────

# (key, header label) per column on the daily sheet — mirrors the page tables
_EXCEL_DAILY_COLUMNS = (
    ("employee",        _("Employee ID")),
    ("employee_name",   _("Employee Name")),
    ("department",      _("Department")),
    ("designation",     _("Designation")),
    ("zk_biometric_device", _("Device")),
    ("attendance_device_id", _("Device ID")),
    ("shift_type",      _("Shift")),
    ("date",            _("Date")),
    ("weekday",         _("Day")),
    ("status",          _("Status")),
    ("hours",           _("Hours")),
    ("formatted",       _("Formatted Hours")),
    ("day_ot_hours",    _("Day OT")),
    ("night_ot_hours",  _("Night OT")),
    ("weekend_ot_hours", _("Weekend OT")),
    ("holiday_ot_hours", _("Holiday OT")),
    ("overtime_hours",  _("Total OT")),
    ("late_minutes",    _("Late (min)")),
    ("early_minutes",   _("Early (min)")),
    ("holiday_name",    _("Holiday Name")),
    ("checkins",        _("Check-ins")),
)

# (key, header label) per column on the summary sheet
_EXCEL_SUMMARY_COLUMNS = (
    ("employee",        _("Employee ID")),
    ("employee_name",   _("Employee Name")),
    ("department",      _("Department")),
    ("designation",     _("Designation")),
    ("zk_biometric_device", _("Device")),
    ("attendance_device_id", _("Device ID")),
    ("shift_type",      _("Shift")),
    ("working_days",    _("Working Days")),
    ("absent_days",     _("Absent Days")),
    ("invalid_days",    _("Invalid Days")),
    ("review_days",     _("Review Needed")),
    ("total_hours",     _("Total Hours")),
    ("day_ot_hours",    _("Day OT")),
    ("night_ot_hours",  _("Night OT")),
    ("weekend_ot_hours", _("Weekend OT")),
    ("holiday_ot_hours", _("Holiday OT")),
    ("overtime_hours",  _("Total OT")),
)

_EXCEL_COLUMN_WIDTHS = {
    "employee": 12, "employee_name": 22, "department": 16, "designation": 16,
    "zk_biometric_device": 12, "attendance_device_id": 12, "shift_type": 14,
    "date": 12, "weekday": 10, "status": 14, "hours": 10, "formatted": 14,
    "day_ot_hours": 10, "night_ot_hours": 10, "weekend_ot_hours": 12,
    "holiday_ot_hours": 12, "overtime_hours": 10, "late_minutes": 10,
    "early_minutes": 10, "holiday_name": 18, "checkins": 46,
    "working_days": 12, "absent_days": 12, "invalid_days": 12,
    "review_days": 12, "total_hours": 10,
}


def _checkins_text(checkins):
    """Render a day's check-ins as readable text, e.g. '08:57 IN; 17:30 OUT'."""
    parts = []
    for c in checkins or []:
        log_type = cstr(c.get("log_type") or "IN").upper()
        text = "{0} {1}".format(cstr(c.get("time") or "")[:5], log_type)
        if c.get("is_overtime"):
            text += " (OT)"
        if c.get("manually_edited"):
            text += " *"
        if c.get("ignored"):
            text += " (ignored)"
        parts.append(text)
    return "; ".join(parts)


def _employee_summary(emp):
    """Aggregate an employee's daily rows into summary numbers.
    Mirrors the attendance processor: Present/Holiday count as working days,
    Half Day as 0.5; Absent / Invalid / Manual Review are counted separately."""
    days = emp.get("days") or []
    summary = {
        "working_days": 0.0,
        "absent_days":  0,
        "invalid_days": 0,
        "review_days":  0,
        "total_hours":  0.0,
        "day_ot_hours":    0.0,
        "night_ot_hours":  0.0,
        "weekend_ot_hours": 0.0,
        "holiday_ot_hours": 0.0,
        "overtime_hours":  0.0,
    }
    for d in days:
        status = cstr(d.get("status") or "")
        if status == "Present":
            summary["working_days"] += 1.0
        elif status == "Half Day":
            summary["working_days"] += 0.5
        elif status == "Absent":
            summary["absent_days"] += 1
        elif status == "Invalid":
            summary["invalid_days"] += 1
        elif status == "Manual Review":
            summary["review_days"] += 1
        elif status == "Holiday":
            summary["working_days"] += 1.0

        summary["total_hours"] += flt(d.get("hours") or 0)
        summary["day_ot_hours"]     += flt(d.get("day_ot_hours") or 0)
        summary["night_ot_hours"]   += flt(d.get("night_ot_hours") or 0)
        summary["weekend_ot_hours"] += flt(d.get("weekend_ot_hours") or 0)
        summary["holiday_ot_hours"] += flt(d.get("holiday_ot_hours") or 0)
        summary["overtime_hours"]   += flt(d.get("overtime_hours") or 0)
    return summary


def _emp_display_name(emp):
    return (emp.get("fullname") or emp.get("employee_name") or emp.get("employee") or "")


def _build_excel_workbook(data):
    """
    Build the Daily Check-ins Excel workbook and return it as bytes.
    Sheet 1 'Daily Checkins' holds one row per employee-day; sheet 2
    'Summary' holds per-employee aggregates plus a grand-total row.
    """
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from frappe.utils import now_datetime

    employees = data.get("employees") or []
    period = "{0} to {1}".format(
        cstr(data.get("from_date") or ""), cstr(data.get("to_date") or ""))

    meta_bits = []
    if data.get("company"):
        meta_bits.append(_("Company: {0}").format(data.get("company")))
    if data.get("attendance_summary"):
        meta_bits.append(_("Attendance Summary: {0}").format(data.get("attendance_summary")))
    meta_bits.append(_("Generated by {0} on {1}").format(
        frappe.session.user, now_datetime().strftime("%d %b %Y, %H:%M")))
    meta_line = "  |  ".join(meta_bits)

    hdr_fill  = PatternFill("solid", fgColor="16324F")
    hdr_font  = Font(color="FFFFFF", bold=True)
    thin      = Side(style="thin", color="D6E2EF")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_font = Font(bold=True, size=13, color="16324F")
    meta_font  = Font(size=9, color="5B6B7A")

    def write_title_block(ws, ncols):
        last_col = get_column_letter(ncols)
        ws.merge_cells("A1:{0}1".format(last_col))
        ws["A1"] = _("Employee Daily Check-ins") + (" — {0}".format(period) if period != " to " else "")
        ws["A1"].font = title_font
        ws.merge_cells("A2:{0}2".format(last_col))
        ws["A2"] = meta_line
        ws["A2"].font = meta_font

    def write_header_row(ws, row_idx, columns):
        for col_idx, (_key, label) in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=label)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx, (key, _label) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = _EXCEL_COLUMN_WIDTHS.get(key, 12)

    wb = Workbook()

    # ── Sheet 1: Daily Checkins ───────────────────────────────────────────
    ws_daily = wb.active
    ws_daily.title = _("Daily Checkins")
    write_title_block(ws_daily, len(_EXCEL_DAILY_COLUMNS))

    hdr_row = 4
    write_header_row(ws_daily, hdr_row, _EXCEL_DAILY_COLUMNS)

    r = hdr_row + 1
    for emp in employees:
        emp_name = _emp_display_name(emp)
        for d in emp.get("days") or []:
            values = {
                "employee":        emp.get("employee") or "",
                "employee_name":   emp_name,
                "department":      emp.get("department") or "",
                "designation":     emp.get("designation") or "",
                "zk_biometric_device": emp.get("zk_biometric_device") or "",
                "attendance_device_id": emp.get("attendance_device_id") or "",
                "shift_type":      emp.get("shift_type") or "",
                "date":            cstr(d.get("date") or ""),
                "weekday":         _(cstr(d.get("weekday") or "")),
                "status":          _(cstr(d.get("status") or "")),
                "hours":           flt(d.get("hours") or 0),
                "formatted":       _hhmm(d.get("hours")),
                "day_ot_hours":    flt(d.get("day_ot_hours") or 0),
                "night_ot_hours":  flt(d.get("night_ot_hours") or 0),
                "weekend_ot_hours": flt(d.get("weekend_ot_hours") or 0),
                "holiday_ot_hours": flt(d.get("holiday_ot_hours") or 0),
                "overtime_hours":  flt(d.get("overtime_hours") or 0),
                "late_minutes":    int(flt(d.get("late_minutes") or 0)),
                "early_minutes":   int(flt(d.get("early_minutes") or 0)),
                "holiday_name":    d.get("holiday_name") or "",
                "checkins":        _checkins_text(d.get("checkins")),
            }
            for col_idx, (key, _label) in enumerate(_EXCEL_DAILY_COLUMNS, start=1):
                cell = ws_daily.cell(row=r, column=col_idx, value=values[key])
                cell.border = border
                if key in ("hours", "day_ot_hours", "night_ot_hours",
                           "weekend_ot_hours", "holiday_ot_hours", "overtime_hours"):
                    cell.number_format = "0.00"
            r += 1

    ws_daily.freeze_panes = "A{0}".format(hdr_row + 1)
    ws_daily.auto_filter.ref = "A{0}:{1}{2}".format(
        hdr_row, get_column_letter(len(_EXCEL_DAILY_COLUMNS)), max(r - 1, hdr_row))

    # ── Sheet 2: Summary ──────────────────────────────────────────────────
    ws_summary = wb.create_sheet(_("Summary"))
    write_title_block(ws_summary, len(_EXCEL_SUMMARY_COLUMNS))
    write_header_row(ws_summary, hdr_row, _EXCEL_SUMMARY_COLUMNS)

    r = hdr_row + 1
    for emp in employees:
        summary = _employee_summary(emp)
        values = {
            "employee":        emp.get("employee") or "",
            "employee_name":   _emp_display_name(emp),
            "department":      emp.get("department") or "",
            "designation":     emp.get("designation") or "",
            "zk_biometric_device": emp.get("zk_biometric_device") or "",
            "attendance_device_id": emp.get("attendance_device_id") or "",
            "shift_type":      emp.get("shift_type") or "",
        }
        values.update(summary)
        for col_idx, (key, _label) in enumerate(_EXCEL_SUMMARY_COLUMNS, start=1):
            cell = ws_summary.cell(row=r, column=col_idx, value=values[key])
            cell.border = border
            if key in ("working_days", "total_hours", "day_ot_hours",
                       "night_ot_hours", "weekend_ot_hours",
                       "holiday_ot_hours", "overtime_hours"):
                cell.number_format = "0.00"
        r += 1

    # Grand totals row
    if employees:
        grand = {
            "working_days": 0.0, "absent_days": 0, "invalid_days": 0,
            "review_days": 0, "total_hours": 0.0, "day_ot_hours": 0.0,
            "night_ot_hours": 0.0, "weekend_ot_hours": 0.0,
            "holiday_ot_hours": 0.0, "overtime_hours": 0.0,
        }
        for emp in employees:
            s = _employee_summary(emp)
            for key in grand:
                grand[key] += s[key]

        grand_fill  = PatternFill("solid", fgColor="E6EDF5")
        grand_font  = Font(bold=True, color="16324F")
        for col_idx, (key, _label) in enumerate(_EXCEL_SUMMARY_COLUMNS, start=1):
            cell = ws_summary.cell(row=r, column=col_idx,
                                   value=_("Grand Total") if col_idx == 1 else grand.get(key))
            cell.fill = grand_fill
            cell.font = grand_font
            cell.border = border
            if col_idx > 1 and key in ("working_days", "total_hours", "day_ot_hours",
                                       "night_ot_hours", "weekend_ot_hours",
                                       "holiday_ot_hours", "overtime_hours"):
                cell.number_format = "0.00"

    ws_summary.freeze_panes = "A{0}".format(hdr_row + 1)
    ws_summary.auto_filter.ref = "A{0}:{1}{2}".format(
        hdr_row, get_column_letter(len(_EXCEL_SUMMARY_COLUMNS)), max(r - 1, hdr_row))

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
